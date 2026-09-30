'''Conciliation reports with pt-BR export.'''

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from src.config import APP_VERSION
from src.logger import get_logger

logger = get_logger(__name__)

_REPORT_HEADERS = (
    'Origem',
    'Data original',
    'Data normalizada',
    'Descrição original',
    'Valor original',
    'Valor normalizado',
    'Sinal',
    'Status',
    'Par ID',
    'Diferença dias',
    'Diferença valor',
    'Score descrição',
    'Regra ID',
    'Motivo',
    'Ação manual',
)

_LOG_HEADERS_PT = (
    'Par ID',
    'Regra ID',
    'Diferença dias',
    'Diferença valor',
    'Score descrição',
    'Motivo',
    'Código erro',
    'Ação manual',
    'Data/Hora ação',
    'Parâmetros',
)

_ERROR_HEADERS_PT = ('Linha', 'Motivo', 'Como corrigir')

_STATUS_PT = {
    'auto': 'Conciliada',
    'potential': 'Para revisão',
    'pending': 'Pendente',
    'divergent': 'Divergente',
    'duplicate': 'Duplicada',
}


def format_brl(value: object) -> str:
    '''Format numeric input as pt-BR currency text.'''
    if value is None:
        return '—'
    try:
        if pd.isna(value):
            return '—'
    except (ValueError, TypeError):
        pass
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return str(value)
    quantized = amount.quantize(Decimal('0.00'))
    text = f'{quantized:,.2f}'
    text = text.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'R$ {text}'


def calc_kpis(results: dict, statement_count: int = 0, ledger_count: int = 0) -> dict:
    '''Calculate totals and percentages from result tables.'''
    auto_count = _safe_len(results.get('auto'))
    potential_count = _safe_len(results.get('potential'))
    pending_count = _safe_len(results.get('pending'))
    divergent_count = _safe_len(results.get('divergent'))
    duplicate_count = _safe_len(results.get('duplicate'))
    unified = auto_count + potential_count + pending_count + divergent_count
    unified = unified + duplicate_count
    if unified == 0:
        return _empty_kpis(statement_count, ledger_count)
    pct_auto = round(auto_count / unified * 100, 1)
    pct_review = round(potential_count / unified * 100, 1)
    pct_pending = round(pending_count / unified * 100, 1)
    pct_divergent = round((divergent_count + duplicate_count) / unified * 100, 1)
    return {
        'total_statement': int(statement_count),
        'total_ledger': int(ledger_count),
        'pct_auto': float(pct_auto),
        'pct_review': float(pct_review),
        'pct_pending': float(pct_pending),
        'pct_divergent': float(pct_divergent),
        'exception_rate': float(round(100.0 - pct_auto, 1)),
    }


def _safe_len(frame: object) -> int:
    '''Return row count handling none.'''
    if frame is None:
        return 0
    try:
        return int(len(frame))
    except TypeError:
        return 0


def _empty_kpis(statement_count: int, ledger_count: int) -> dict:
    '''Build zeroed kpi dict.'''
    return {
        'total_statement': int(statement_count),
        'total_ledger': int(ledger_count),
        'pct_auto': 0.0,
        'pct_review': 0.0,
        'pct_pending': 0.0,
        'pct_divergent': 0.0,
        'exception_rate': 100.0,
    }


def _params_snapshot(params: object) -> dict:
    '''Build deterministic snapshot with app version.'''
    base = {} if not isinstance(params, dict) else dict(params)
    tolerance = base.get('value_tolerance', Decimal('0.00'))
    version = base.get('APP_VERSION', base.get('app_version', APP_VERSION))
    return {
        'date_tolerance_days': base.get('date_tolerance_days', 2),
        'fuzzy_threshold': base.get('fuzzy_threshold', 85),
        'value_tolerance': tolerance,
        'use_fuzzy': base.get('use_fuzzy', True),
        'APP_VERSION': str(version),
    }


def _snapshot_text(snapshot: dict) -> str:
    '''Render snapshot as single audit string.'''
    day = snapshot.get('date_tolerance_days')
    fuzzy = snapshot.get('fuzzy_threshold')
    tolerance = snapshot.get('value_tolerance')
    version = snapshot.get('APP_VERSION', snapshot.get('app_version', APP_VERSION))
    first = f'date_tolerance_days={day}; fuzzy_threshold={fuzzy}'
    second = f'value_tolerance={tolerance}; APP_VERSION={version}'
    return f'{first}; {second}'


def _manual_lookup(review_log: object) -> dict:
    '''Index manual decisions by match id keeping last.'''
    lookup: dict = {}
    if not isinstance(review_log, list):
        return lookup
    for entry in review_log:
        if not isinstance(entry, dict):
            continue
        pair = entry.get('match_id')
        if pair is None:
            continue
        lookup[str(pair)] = {
            'review_action': entry.get('review_action', ''),
            'timestamp': entry.get('timestamp', ''),
        }
    return lookup


def _manual_label(action: object) -> str:
    '''Translate review action to pt-BR label.'''
    text = str(action) if action is not None else ''
    if text == 'conciliada_manual':
        return 'Confirmada manualmente'
    if text == 'rejeitada':
        return 'Rejeitada'
    if text == 'Confirmada manualmente':
        return text
    if text == 'Rejeitada':
        return text
    return ''


def _status_label(category: str) -> str:
    '''Map internal category to pt-BR status.'''
    return _STATUS_PT.get(str(category), str(category))


def _reason_label(raw: object) -> str:
    '''Translate internal motive code to pt-BR text.'''
    labels = {
        'regra_composta_ok': 'Valor, data e descrição conferem',
        'descricao_baixa_similaridade': 'Descrição com baixa similaridade',
        'descricao_divergente': 'Descrição divergente',
        'divergencia_centavos': 'Divergência de centavos',
        'divergencia_valor_descricao': 'Valor e descrição divergentes',
        'ambiguidade_multipla': 'Múltiplos candidatos possíveis',
        'sem_candidato': 'Sem candidato próximo',
        'data_fora_tolerancia': 'Data fora da tolerância',
        'valor_fora_tolerancia': 'Valor fora da tolerância',
        'sinal_bloqueado': 'Sinal oposto (débito x crédito)',
        'duplicada_suspeita': 'Duplicada suspeita',
    }
    return labels.get(str(raw), str(raw))


def _lookup_field(frame: object, idx: object, fields: list) -> object:
    '''Return first available field value for index.'''
    if frame is None or idx is None:
        return None
    try:
        if len(frame) == 0 or idx not in frame.index:
            return None
        row = frame.loc[idx]
    except (KeyError, ValueError, TypeError):
        return None
    for field in fields:
        try:
            if field not in frame.columns:
                continue
            found = row[field]
        except (KeyError, ValueError):
            continue
        if found is None:
            continue
        try:
            if pd.isna(found):
                continue
        except (ValueError, TypeError):
            pass
        return found
    return None


def _format_iso(value: object) -> str:
    '''Format date like value as YYYY-MM-DD string.'''
    if value is None:
        return ''
    try:
        if pd.isna(value):
            return ''
    except (ValueError, TypeError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.date().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text == '' or text.lower() in ('nan', 'nat', 'none'):
        return ''
    return text


def _as_text(value: object) -> str:
    '''Convert cell value to stable text.'''
    if value is None:
        return ''
    try:
        if pd.isna(value):
            return ''
    except (ValueError, TypeError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, datetime):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, date):
        return value.strftime('%d/%m/%Y')
    return str(value)


def _as_number(value: object) -> object:
    '''Convert Decimal numeric to float preserving none.'''
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (ValueError, TypeError):
        pass
    if isinstance(value, bool):
        return None
    try:
        return float(Decimal(str(value)))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _extract_side(frame: object, idx: object) -> dict:
    '''Extract original normalized triple for index.'''
    orig_date = _lookup_field(frame, idx, ['Data', 'event_date'])
    norm_date = _lookup_field(frame, idx, ['normalized_date'])
    orig_desc = _lookup_field(frame, idx, ['Descrição', 'description'])
    orig_amount = _lookup_field(frame, idx, ['Valor', 'amount'])
    norm_value = _lookup_field(frame, idx, ['normalized_value', 'normalized_amount'])
    sign = _lookup_field(frame, idx, ['sign'])
    return {
        'orig_date': _as_text(orig_date),
        'norm_date': _format_iso(norm_date),
        'orig_desc': _as_text(orig_desc),
        'orig_amount': _as_text(orig_amount),
        'norm_value': _as_number(norm_value),
        'sign': _as_number(sign),
    }


def _matched_rows(
    match_frame: object, statement_df: object, ledger_frame: object, status: str, manual: dict
) -> list:
    '''Expand matched pairs into two pt-BR rows.'''
    rows: list = []
    if match_frame is None or len(match_frame) == 0:
        return rows
    for _, item in match_frame.iterrows():
        pair = item.get('match_id')
        manual_entry = manual.get(str(pair), {})
        label = _manual_label(manual_entry.get('review_action', ''))
        rows.append(_matched_side(item, statement_df, True, status, label))
        rows.append(_matched_side(item, ledger_frame, False, status, label))
    return rows


def _matched_side(item: object, frame: object, is_statement: bool, status: str, label: str) -> dict:
    '''Build single side row for matched pair.'''
    if is_statement:
        idx = item.get('statement_idx')
        origin = 'Extrato'
    else:
        idx = item.get('ledger_idx')
        origin = 'Interno'
    side = _extract_side(frame, idx)
    reason = item.get('reason')
    motive = '' if reason is None else _reason_label(reason)
    return {
        'Origem': origin,
        'Data original': side['orig_date'],
        'Data normalizada': side['norm_date'],
        'Descrição original': side['orig_desc'],
        'Valor original': side['orig_amount'],
        'Valor normalizado': side['norm_value'],
        'Sinal': side['sign'],
        'Status': status,
        'Par ID': item.get('match_id') if item.get('match_id') is not None else '',
        'Diferença dias': item.get('day_diff'),
        'Diferença valor': _as_number(item.get('value_diff')),
        'Score descrição': item.get('description_score'),
        'Regra ID': item.get('rule_id') if item.get('rule_id') is not None else '',
        'Motivo': motive,
        'Ação manual': label,
    }


def _single_side_rows(
    match_frame: object, statement_df: object, ledger_frame: object, status: str, manual: dict
) -> list:
    '''Build one row per pending duplicate entry.'''
    rows: list = []
    if match_frame is None or len(match_frame) == 0:
        return rows
    for _, item in match_frame.iterrows():
        rows.append(_single_side_row(item, statement_df, ledger_frame, status, manual))
    return rows


def _resolve_side_frame(item: object, statement_df: object, ledger_frame: object) -> tuple:
    '''Resolve index frame origin for single side row.'''
    source = item.get('source')
    if source == 'ledger':
        return (item.get('ledger_idx'), ledger_frame, 'Interno')
    if source == 'statement':
        return (item.get('statement_idx'), statement_df, 'Extrato')
    if item.get('ledger_idx') is not None and source not in ('statement', 'ledger'):
        return (item.get('ledger_idx'), ledger_frame, 'Interno')
    return (item.get('statement_idx'), statement_df, 'Extrato')


def _single_side_row(
    item: object, statement_df: object, ledger_frame: object, status: str, manual: dict
) -> dict:
    '''Build pending duplicate row resolving side.'''
    idx, frame, origin = _resolve_side_frame(item, statement_df, ledger_frame)
    side = _extract_side(frame, idx)
    pair = item.get('match_id')
    manual_entry = manual.get(str(pair), {}) if pair is not None else {}
    return _single_row_dict(item, side, origin, status, pair, manual_entry)


def _single_row_dict(
    item: object, side: dict, origin: str, status: str, pair: object, manual_entry: dict
) -> dict:
    '''Build dict for single side detail row.'''
    reason = item.get('reason')
    motive = '' if reason is None else _reason_label(reason)
    return {
        'Origem': origin,
        'Data original': side['orig_date'],
        'Data normalizada': side['norm_date'],
        'Descrição original': side['orig_desc'],
        'Valor original': side['orig_amount'],
        'Valor normalizado': side['norm_value'],
        'Sinal': side['sign'],
        'Status': status,
        'Par ID': pair if pair is not None else '',
        'Diferença dias': item.get('day_diff'),
        'Diferença valor': _as_number(item.get('value_diff')),
        'Score descrição': item.get('description_score'),
        'Regra ID': item.get('rule_id') if item.get('rule_id') is not None else '',
        'Motivo': motive,
        'Ação manual': _manual_label(manual_entry.get('review_action', '')),
    }


def _normalize_error_frames(error_frames: object) -> pd.DataFrame:
    '''Consolidate error inputs into pt-BR display frame.'''
    if error_frames is None:
        return pd.DataFrame(columns=list(_ERROR_HEADERS_PT))
    if isinstance(error_frames, dict):
        frames = list(error_frames.values())
        return _concat_error_frames(frames)
    if isinstance(error_frames, (list, tuple)):
        return _concat_error_frames(list(error_frames))
    if isinstance(error_frames, pd.DataFrame):
        return _clean_error_frame(error_frames)
    return pd.DataFrame(columns=list(_ERROR_HEADERS_PT))


def _concat_error_frames(frames: list) -> pd.DataFrame:
    '''Concat multiple error frames normalizing headers.'''
    cleaned: list = []
    for frame in frames:
        if frame is None or len(frame) == 0:
            continue
        cleaned.append(_clean_error_frame(frame))
    if not cleaned:
        return pd.DataFrame(columns=list(_ERROR_HEADERS_PT))
    merged = pd.concat(cleaned, ignore_index=True)
    return merged[list(_ERROR_HEADERS_PT)]


def _clean_error_frame(frame: pd.DataFrame) -> pd.DataFrame:
    '''Rename Orientação to Como corrigir preserving content.'''
    work = frame.copy()
    if 'Orientação' in list(work.columns) and 'Como corrigir' not in list(work.columns):
        work = work.rename(columns={'Orientação': 'Como corrigir'})
    for field in _ERROR_HEADERS_PT:
        if field not in list(work.columns):
            work[field] = ''
    return work[list(_ERROR_HEADERS_PT)]


def _error_detail_rows(error_frames: object) -> list:
    '''Convert error summaries to exception detail rows.'''
    rows: list = []
    if error_frames is None:
        return rows
    if isinstance(error_frames, dict):
        for key, frame in error_frames.items():
            origin = 'Interno' if str(key).lower().startswith('ledger') else 'Extrato'
            rows.extend(_error_frame_rows(frame, origin))
        return rows
    if isinstance(error_frames, (list, tuple)):
        for pos, frame in enumerate(error_frames):
            origin = 'Interno' if pos == 1 and len(error_frames) == 2 else 'Extrato'
            rows.extend(_error_frame_rows(frame, origin))
        return rows
    if isinstance(error_frames, pd.DataFrame):
        return _error_frame_rows(error_frames, 'Extrato')
    return rows


def _error_frame_rows(frame: object, origin: str) -> list:
    '''Build detail rows for single error frame.'''
    rows: list = []
    if frame is None or len(frame) == 0:
        return rows
    clean = _clean_error_frame(frame)
    for _, item in clean.iterrows():
        rows.append(
            {
                'Origem': origin,
                'Data original': '',
                'Data normalizada': '',
                'Descrição original': '',
                'Valor original': '',
                'Valor normalizado': None,
                'Sinal': None,
                'Status': 'Erro',
                'Par ID': '',
                'Diferença dias': None,
                'Diferença valor': None,
                'Score descrição': None,
                'Regra ID': '',
                'Motivo': str(item.get('Motivo', '')),
                'Ação manual': '',
            }
        )
    return rows


def build_rule_log(
    results: dict, review_log: object = None, error_frames: object = None
) -> pd.DataFrame:
    '''Build auditable log with rule manual and error fields.'''
    manual = _manual_lookup(review_log)
    snapshot = _params_snapshot(_infer_params(results))
    rows = _collect_match_log_rows(results, manual)
    rows.extend(_collect_error_log_rows(error_frames, snapshot, manual))
    columns = [
        'match_id',
        'rule_id',
        'day_diff',
        'value_diff',
        'description_score',
        'reason',
        'error_code',
        'review_action',
        'timestamp',
        'params_snapshot',
    ]
    if not rows:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame(rows, columns=columns)


def _infer_params(results: dict) -> dict:
    '''Infer params snapshot from first available row.'''
    if not isinstance(results, dict):
        return {}
    for key in ('auto', 'potential', 'pending', 'divergent', 'duplicate'):
        frame = results.get(key)
        if frame is not None and len(frame) > 0 and 'params_snapshot' in list(frame.columns):
            first = frame.iloc[0]['params_snapshot']
            if isinstance(first, dict):
                return dict(first)
    return {}


def _collect_match_log_rows(results: dict, manual: dict) -> list:
    '''Collect log rows for five result tables.'''
    rows: list = []
    if not isinstance(results, dict):
        return rows
    for key in ('auto', 'potential', 'pending', 'divergent', 'duplicate'):
        frame = results.get(key)
        if frame is None or len(frame) == 0:
            continue
        for _, item in frame.iterrows():
            rows.append(_match_log_row(item, manual))
    return rows


def _match_log_row(item: object, manual: dict) -> dict:
    '''Build single log row with manual overlay.'''
    pair = item.get('match_id')
    entry = manual.get(str(pair), {}) if pair is not None else {}
    snapshot = item.get('params_snapshot')
    clean_snapshot = dict(snapshot) if isinstance(snapshot, dict) else {}
    return {
        'match_id': pair,
        'rule_id': item.get('rule_id'),
        'day_diff': item.get('day_diff'),
        'value_diff': item.get('value_diff'),
        'description_score': item.get('description_score'),
        'reason': item.get('reason'),
        'error_code': '',
        'review_action': entry.get('review_action', ''),
        'timestamp': entry.get('timestamp', ''),
        'params_snapshot': clean_snapshot,
    }


def _collect_error_log_rows(error_frames: object, snapshot: dict, manual: dict) -> list:
    '''Collect log rows for validation errors.'''
    rows: list = []
    display = _normalize_error_frames(error_frames)
    if len(display) == 0:
        return rows
    for _, item in display.iterrows():
        rows.append(_error_log_row(item, snapshot))
    return rows


def _error_log_row(item: object, snapshot: dict) -> dict:
    '''Build log row inferring error code from motive.'''
    motive = str(item.get('Motivo', ''))
    code = _infer_error_code(motive)
    return {
        'match_id': None,
        'rule_id': 'RN-09',
        'day_diff': None,
        'value_diff': None,
        'description_score': None,
        'reason': motive,
        'error_code': code,
        'review_action': '',
        'timestamp': '',
        'params_snapshot': dict(snapshot),
    }


def _infer_error_code(motive: str) -> str:
    '''Infer machine code from pt-BR motive text.'''
    text = str(motive).lower()
    if 'data inválida' in text or 'data invalida' in text:
        return 'DATA_INVALIDA'
    if 'valor inválido' in text or 'valor invalido' in text:
        return 'VALOR_INVALIDO'
    if 'coluna obrigatória' in text or 'coluna obrigatoria' in text:
        return 'COLUNA_AUSENTE'
    return 'ERRO_VALIDACAO'


def _kpi_labels(kpis: dict) -> list:
    '''Build pt-BR kpi label value pairs.'''
    return [
        ('Total extrato', kpis.get('total_statement', 0)),
        ('Total interno', kpis.get('total_ledger', 0)),
        ('% Conciliado', kpis.get('pct_auto', 0.0)),
        ('% Para revisão', kpis.get('pct_review', 0.0)),
        ('% Pendente', kpis.get('pct_pending', 0.0)),
        ('% Divergente', kpis.get('pct_divergent', 0.0)),
        ('Taxa de exceção', kpis.get('exception_rate', 0.0)),
    ]


def _param_rows(snapshot: dict) -> list:
    '''Build snapshot key value pairs.'''
    version = snapshot.get('APP_VERSION', snapshot.get('app_version'))
    return [
        ('date_tolerance_days', snapshot.get('date_tolerance_days')),
        ('fuzzy_threshold', snapshot.get('fuzzy_threshold')),
        ('value_tolerance', str(snapshot.get('value_tolerance'))),
        ('APP_VERSION', version),
    ]


def _write_resumo(sheet: object, kpis: dict, snapshot: dict) -> None:
    '''Write kpi and snapshot blocks into resumo sheet.'''
    sheet['A1'] = 'Resumo da conciliação'
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A3'] = 'Indicador'
    sheet['B3'] = 'Valor'
    row_no = _write_kpi_block(sheet, kpis, 4)
    _write_param_block(sheet, snapshot, row_no)


def _write_kpi_block(sheet: object, kpis: dict, start: int) -> int:
    '''Write kpi rows returning next free row.'''
    row_no = int(start)
    for label, value in _kpi_labels(kpis):
        sheet.cell(row=row_no, column=1, value=label)
        sheet.cell(row=row_no, column=2, value=value)
        row_no = row_no + 1
    return row_no


def _write_param_block(sheet: object, snapshot: dict, row_no: int) -> None:
    '''Write snapshot rows after kpi block.'''
    sheet.cell(row=row_no + 1, column=1, value='Parâmetro')
    sheet.cell(row=row_no + 1, column=2, value='Valor')
    line = row_no + 2
    for key, value in _param_rows(snapshot):
        sheet.cell(row=line, column=1, value=key)
        sheet.cell(row=line, column=2, value=value)
        line = line + 1


def _write_detail(book: Workbook, title: str, rows: list) -> None:
    '''Create detail sheet with pt-BR headers and rows.'''
    sheet = book.create_sheet(title=title)
    for pos, header in enumerate(_REPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    line = 2
    for item in rows:
        for pos, header in enumerate(_REPORT_HEADERS, start=1):
            sheet.cell(row=line, column=pos, value=item.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _write_errors(book: Workbook, error_frames: object) -> None:
    '''Create erros sheet with pt-BR correction headers.'''
    sheet = book.create_sheet(title='Erros')
    for pos, header in enumerate(_ERROR_HEADERS_PT, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    display = _normalize_error_frames(error_frames)
    line = 2
    for _, item in display.iterrows():
        for pos, header in enumerate(_ERROR_HEADERS_PT, start=1):
            sheet.cell(row=line, column=pos, value=item.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _write_log(book: Workbook, log_frame: pd.DataFrame, snapshot: dict) -> None:
    '''Create log sheet translating internal fields to pt-BR.'''
    sheet = book.create_sheet(title='Log_Regras')
    for pos, header in enumerate(_LOG_HEADERS_PT, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    text = _snapshot_text(snapshot)
    line = 2
    if log_frame is None or len(log_frame) == 0:
        return
    for _, item in log_frame.iterrows():
        values = _log_display_values(item, text)
        for pos, header in enumerate(_LOG_HEADERS_PT, start=1):
            sheet.cell(row=line, column=pos, value=values.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _log_display_values(item: object, text: str) -> dict:
    '''Map internal log row to pt-BR display values.'''
    snapshot = item.get('params_snapshot')
    if isinstance(snapshot, dict) and len(snapshot) > 0:
        detail = _snapshot_text(_params_snapshot(snapshot))
    else:
        detail = text
    reason = item.get('reason')
    motive = '' if reason is None else _reason_label(reason)
    return {
        'Par ID': item.get('match_id') if item.get('match_id') is not None else '',
        'Regra ID': item.get('rule_id') if item.get('rule_id') is not None else '',
        'Diferença dias': item.get('day_diff'),
        'Diferença valor': _as_number(item.get('value_diff')),
        'Score descrição': item.get('description_score'),
        'Motivo': motive,
        'Código erro': item.get('error_code') if item.get('error_code') is not None else '',
        'Ação manual': _manual_label(item.get('review_action', '')),
        'Data/Hora ação': item.get('timestamp') if item.get('timestamp') is not None else '',
        'Parâmetros': detail,
    }


def _collect_all_detail(
    results: dict, statement_df: object, ledger_frame: object, manual: dict
) -> dict:
    '''Collect detail rows for five categories.'''
    auto_rows = _matched_rows(
        results.get('auto'), statement_df, ledger_frame, _status_label('auto'), manual
    )
    potential_rows = _matched_rows(
        results.get('potential'), statement_df, ledger_frame, _status_label('potential'), manual
    )
    pending_rows = _single_side_rows(
        results.get('pending'), statement_df, ledger_frame, _status_label('pending'), manual
    )
    divergent_rows = _matched_rows(
        results.get('divergent'), statement_df, ledger_frame, _status_label('divergent'), manual
    )
    duplicate_rows = _single_side_rows(
        results.get('duplicate'), statement_df, ledger_frame, _status_label('duplicate'), manual
    )
    return {
        'auto': auto_rows,
        'potential': potential_rows,
        'pending': pending_rows,
        'divergent': divergent_rows,
        'duplicate': duplicate_rows,
    }


def _write_all_details(book: Workbook, detail: dict) -> None:
    '''Write five detail sheets into workbook.'''
    _write_detail(book, 'Conciliadas', detail.get('auto', []))
    _write_detail(book, 'Para_Revisao', detail.get('potential', []))
    _write_detail(book, 'Pendentes', detail.get('pending', []))
    _write_detail(book, 'Divergentes', detail.get('divergent', []))
    _write_detail(book, 'Duplicadas', detail.get('duplicate', []))


def build_conciliation_workbook(
    results: dict,
    params: dict,
    statement_df: object = None,
    ledger_frame: object = None,
    error_frames: object = None,
    review_log: object = None,
) -> bytes:
    '''Generate xlsx bytes with resumo details errors and log.'''
    snapshot = _params_snapshot(params)
    statement_count = 0 if statement_df is None else _safe_len(statement_df)
    ledger_count = 0 if ledger_frame is None else _safe_len(ledger_frame)
    kpis = calc_kpis(results, statement_count, ledger_count)
    manual = _manual_lookup(review_log)
    detail = _collect_all_detail(results, statement_df, ledger_frame, manual)
    log_frame = build_rule_log(results, review_log, error_frames)
    book = Workbook()
    book.remove(book.active)
    resumo = book.create_sheet(title='Resumo')
    _write_resumo(resumo, kpis, snapshot)
    _write_all_details(book, detail)
    _write_errors(book, error_frames)
    _write_log(book, log_frame, snapshot)
    logger.info('Built conciliation workbook with snapshot version.')
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _exceptions_frame(
    results: dict,
    statement_df: object,
    ledger_frame: object,
    error_frames: object,
    manual: dict,
) -> pd.DataFrame:
    '''Build sorted exceptions frame excluding auto.'''
    detail = _collect_all_detail(results, statement_df, ledger_frame, manual)
    rows: list = []
    rows.extend(detail.get('potential', []))
    rows.extend(detail.get('pending', []))
    rows.extend(detail.get('divergent', []))
    rows.extend(detail.get('duplicate', []))
    rows.extend(_error_detail_rows(error_frames))
    if not rows:
        return pd.DataFrame(columns=list(_REPORT_HEADERS))
    frame = pd.DataFrame(rows, columns=list(_REPORT_HEADERS))
    frame['_sort_key'] = frame['Valor normalizado'].map(_sort_number)
    frame = frame.sort_values(by='_sort_key', ascending=False, na_position='last')
    return frame.drop(columns=['_sort_key']).reset_index(drop=True)


def _sort_number(value: object) -> object:
    '''Convert valor to sortable float keeping none last.'''
    number = _as_number(value)
    if number is None:
        return float('-inf')
    return float(number)


def build_exceptions_csv(
    results: dict,
    params: object = None,
    statement_df: object = None,
    ledger_frame: object = None,
    error_frames: object = None,
    review_log: object = None,
) -> bytes:
    '''Generate csv bytes with only exceptions ordered desc.'''
    snapshot = _params_snapshot(params if isinstance(params, dict) else _infer_params(results))
    manual = _manual_lookup(review_log)
    frame = _exceptions_frame(results, statement_df, ledger_frame, error_frames, manual)
    buffer = io.StringIO()
    day = snapshot.get('date_tolerance_days')
    fuzzy = snapshot.get('fuzzy_threshold')
    tolerance = snapshot.get('value_tolerance')
    version = snapshot.get('APP_VERSION', snapshot.get('app_version'))
    buffer.write(f'# date_tolerance_days={day}\n')
    buffer.write(f'# fuzzy_threshold={fuzzy}\n')
    buffer.write(f'# value_tolerance={tolerance}\n')
    buffer.write(f'# APP_VERSION={version}\n')
    frame.to_csv(buffer, index=False)
    logger.info('Built exceptions csv with snapshot version.')
    return buffer.getvalue().encode('utf-8')
