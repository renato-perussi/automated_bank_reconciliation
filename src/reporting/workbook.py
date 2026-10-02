'''Workbook sheets with pt-BR headers.'''

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from src.guards import is_missing
from src.logger import get_logger
from src.reporting.details import _collect_all_detail
from src.reporting.determinism import _new_deterministic_workbook, _save_deterministic_bytes
from src.reporting.display_guards import blank_metric, blank_text
from src.reporting.errors import _normalize_error_frames
from src.reporting.formatting import _as_number
from src.reporting.headers import (
    _ERROR_HEADERS_PT,
    _LOG_HEADERS_PT,
    _REPORT_HEADERS,
    _manual_label,
    _manual_lookup,
    _reason_label,
)
from src.reporting.kpis import _safe_len, calc_kpis
from src.reporting.rule_log import build_rule_log
from src.reporting.snapshots import _params_snapshot, _snapshot_text
from src.reporting.summary import (
    _kpi_labels,
    _param_rows,
    _write_kpi_block,
    _write_param_block,
    _write_resumo,
)

logger = get_logger(__name__)

_blank_log_text = blank_text
_blank_log_metric = blank_metric

__all__ = [
    '_kpi_labels',
    '_param_rows',
    '_write_kpi_block',
    '_write_param_block',
    '_write_resumo',
]


def _write_detail(book: Workbook, title: str, rows: list) -> None:
    '''Create detail sheet with pt-BR headers and rows.'''
    sheet = book.create_sheet(title=title)
    for pos, header in enumerate(_REPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    line = 2
    for item in rows:
        for pos, header in enumerate(_REPORT_HEADERS, start=1):
            sheet.cell(row=line, column=pos, value=item.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _write_errors(book: Workbook, error_frames: object) -> None:
    '''Create erros sheet with pt-BR correction headers.'''
    sheet = book.create_sheet(title='Erros')
    for pos, header in enumerate(_ERROR_HEADERS_PT, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    display = _normalize_error_frames(error_frames)
    line = 2
    for _, item in display.iterrows():
        for pos, header in enumerate(_ERROR_HEADERS_PT, start=1):
            sheet.cell(row=line, column=pos, value=item.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _write_log(book: Workbook, log_frame: pd.DataFrame, snapshot: dict) -> None:
    '''Create log sheet translating internal fields to pt-BR.'''
    sheet = book.create_sheet(title='Log_Regras')
    for pos, header in enumerate(_LOG_HEADERS_PT, start=1):
        cell = sheet.cell(row=1, column=pos, value=header)
        cell.font = Font(bold=True)
    text = _snapshot_text(snapshot)
    line = 2
    if log_frame is None or len(log_frame) == 0:
        return
    for _, item in log_frame.iterrows():
        values = _log_display_values(item, text)
        for pos, header in enumerate(_LOG_HEADERS_PT, start=1):
            sheet.cell(row=line, column=pos, value=values.get(header))
        line = line + 1
    sheet.freeze_panes = 'A2'


def _log_display_values(item: object, text: str) -> dict:
    '''Map internal log row to pt-BR display values.'''
    snapshot = item.get('params_snapshot')
    if isinstance(snapshot, dict) and len(snapshot) > 0:
        detail = _snapshot_text(_params_snapshot(snapshot))
    else:
        detail = text
    reason = item.get('reason')
    if reason is None or is_missing(reason):
        motive = ''
    else:
        motive = _reason_label(reason)
    return {
        'Par ID': blank_text(item.get('match_id')),
        'Regra ID': blank_text(item.get('rule_id')),
        'Diferença dias': blank_metric(item.get('day_diff')),
        'Diferença valor': _as_number(item.get('value_diff')),
        'Score descrição': blank_metric(item.get('description_score')),
        'Motivo': motive,
        'Código erro': blank_text(item.get('error_code')),
        'Ação manual': _manual_label(item.get('review_action', '')),
        'Data/Hora ação': blank_text(item.get('timestamp')),
        'Parâmetros': detail,
    }


def _write_all_details(book: Workbook, detail: dict) -> None:
    '''Write five detail sheets into workbook.'''
    _write_detail(book, 'Conciliadas', detail.get('auto', []))
    _write_detail(book, 'Para_Revisao', detail.get('potential', []))
    _write_detail(book, 'Pendentes', detail.get('pending', []))
    _write_detail(book, 'Divergentes', detail.get('divergent', []))
    _write_detail(book, 'Duplicadas', detail.get('duplicate', []))


def build_conciliation_workbook(
    results: dict,
    params: dict,
    statement_df: object = None,
    ledger_frame: object = None,
    error_frames: object = None,
    review_log: object = None,
) -> bytes:
    '''Generate xlsx bytes with resumo details errors and log.'''
    snapshot = _params_snapshot(params)
    statement_count = 0 if statement_df is None else _safe_len(statement_df)
    ledger_count = 0 if ledger_frame is None else _safe_len(ledger_frame)
    kpis = calc_kpis(results, statement_count, ledger_count)
    manual = _manual_lookup(review_log)
    detail = _collect_all_detail(results, statement_df, ledger_frame, manual)
    log_frame = build_rule_log(results, review_log, error_frames)
    book = _new_deterministic_workbook()
    book.remove(book.active)
    resumo = book.create_sheet(title='Resumo')
    _write_resumo(resumo, kpis, snapshot)
    _write_all_details(book, detail)
    _write_errors(book, error_frames)
    _write_log(book, log_frame, snapshot)
    logger.info('Built conciliation workbook with snapshot version.')
    return _save_deterministic_bytes(book)
