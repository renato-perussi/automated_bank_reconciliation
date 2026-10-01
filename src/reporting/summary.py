'''Resumo sheet with kpis and snapshot.'''

from openpyxl.styles import Font


def _kpi_labels(kpis: dict) -> list:
    '''Build pt-BR kpi label value pairs.'''
    return [
        ('Total extrato', kpis.get('total_statement', 0)),
        ('Total interno', kpis.get('total_ledger', 0)),
        ('% Conciliado', kpis.get('pct_auto', 0.0)),
        ('% Para revisão', kpis.get('pct_review', 0.0)),
        ('% Pendente', kpis.get('pct_pending', 0.0)),
        ('% Divergente', kpis.get('pct_divergent', 0.0)),
        ('Taxa de exceção', kpis.get('exception_rate', 0.0)),
    ]


def _param_rows(snapshot: dict) -> list:
    '''Build snapshot key value pairs.'''
    version = snapshot.get('APP_VERSION', snapshot.get('app_version'))
    return [
        ('date_tolerance_days', snapshot.get('date_tolerance_days')),
        ('fuzzy_threshold', snapshot.get('fuzzy_threshold')),
        ('value_tolerance', str(snapshot.get('value_tolerance'))),
        ('APP_VERSION', version),
    ]


def _write_resumo(sheet: object, kpis: dict, snapshot: dict) -> None:
    '''Write kpi and snapshot blocks into resumo sheet.'''
    sheet['A1'] = 'Resumo da conciliação'
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A3'] = 'Indicador'
    sheet['B3'] = 'Valor'
    row_no = _write_kpi_block(sheet, kpis, 4)
    _write_param_block(sheet, snapshot, row_no)


def _write_kpi_block(sheet: object, kpis: dict, start: int) -> int:
    '''Write kpi rows returning next free row.'''
    row_no = int(start)
    for label, value in _kpi_labels(kpis):
        sheet.cell(row=row_no, column=1, value=label)
        sheet.cell(row=row_no, column=2, value=value)
        row_no = row_no + 1
    return row_no


def _write_param_block(sheet: object, snapshot: dict, row_no: int) -> None:
    '''Write snapshot rows after kpi block.'''
    sheet.cell(row=row_no + 1, column=1, value='Parâmetro')
    sheet.cell(row=row_no + 1, column=2, value='Valor')
    line = row_no + 2
    for key, value in _param_rows(snapshot):
        sheet.cell(row=line, column=1, value=key)
        sheet.cell(row=line, column=2, value=value)
        line = line + 1
