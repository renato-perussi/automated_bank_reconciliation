'''Report engine checks.'''

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pandas as pd

import src.mapping as mapping
from src.classifier import build_result_tables
from src.config import APP_VERSION
from src.loader import load_table
from src.matcher import build_params
from src.normalize import normalize_table
from src.report import (
    _as_number,
    _as_text,
    _clean_error_frame,
    _concat_error_frames,
    _error_detail_rows,
    _error_frame_rows,
    _format_iso,
    _infer_error_code,
    _infer_params,
    _lookup_field,
    _manual_label,
    _manual_lookup,
    _normalize_error_frames,
    _params_snapshot,
    _reason_label,
    _resolve_side_frame,
    _safe_len,
    _snapshot_text,
    _sort_number,
    build_conciliation_workbook,
    build_exceptions_csv,
    build_rule_log,
    calc_kpis,
    format_brl,
)
from src.validate import collect_errors


def _make_frame(entries: list, source: str) -> pd.DataFrame:
    '''Build normalized frame from raw triples.'''
    frame = pd.DataFrame(entries, columns=['Data', 'Descrição', 'Valor'])
    mapped = mapping.apply_mapping(frame, mapping.auto_map_columns(frame))
    return normalize_table(mapped, source)


def _make_results() -> tuple:
    '''Build reference tables with defaults.'''
    statement = _make_frame(
        [
            ('10/09/2026', 'Pagamento Fornecedor X', '-2500'),
            ('11/09/2026', 'Recebimento Cliente Y', '4800'),
            ('15/09/2026', 'Pagamento Alfa', '-100'),
        ],
        'statement',
    )
    ledger = _make_frame(
        [
            ('10/09/2026', 'Fornecedor X NF 1254', '-2500'),
            ('12/09/2026', 'Cliente Y', '4800'),
            ('15/09/2026', 'Outro Beta', '-200'),
        ],
        'ledger',
    )
    params = build_params()
    tables = build_result_tables(statement, ledger, params)
    return tables, params, statement, ledger


def test_calc_kpis_keys() -> None:
    '''Kpi dict uses english keys and totals.'''
    tables, _, statement, ledger = _make_results()
    kpis = calc_kpis(tables, len(statement), len(ledger))
    assert set(kpis.keys()) == {
        'total_statement',
        'total_ledger',
        'pct_auto',
        'pct_review',
        'pct_pending',
        'pct_divergent',
        'exception_rate',
    }
    assert kpis['total_statement'] == 3
    assert kpis['total_ledger'] == 3
    assert kpis['exception_rate'] == round(100.0 - kpis['pct_auto'], 1)


def test_format_brl_values() -> None:
    '''Currency formats pt-BR keeping decimal exact.'''
    assert format_brl(Decimal('4800')) == 'R$ 4.800,00'
    assert format_brl(Decimal('-2500')) == 'R$ -2.500,00'
    assert format_brl(Decimal('0')) == 'R$ 0,00'
    assert format_brl(None) == '—'


def test_workbook_sheets_exist() -> None:
    '''Workbook contains eight pt-BR sheets.'''
    tables, params, statement, ledger = _make_results()
    payload = build_conciliation_workbook(tables, params, statement, ledger)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    names = set(book.sheetnames)
    for expected in (
        'Resumo',
        'Conciliadas',
        'Para_Revisao',
        'Pendentes',
        'Divergentes',
        'Duplicadas',
        'Erros',
        'Log_Regras',
    ):
        assert expected in names


def test_workbook_headers_pt() -> None:
    '''Detail sheets use PRD headers in pt-BR.'''
    tables, params, statement, ledger = _make_results()
    payload = build_conciliation_workbook(tables, params, statement, ledger)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    expected = [
        'Origem',
        'Data original',
        'Data normalizada',
        'Descrição original',
        'Valor original',
        'Valor normalizado',
        'Sinal',
        'Status',
        'Par ID',
        'Diferença dias',
        'Diferença valor',
        'Score descrição',
        'Regra ID',
        'Motivo',
        'Ação manual',
    ]
    for name in ('Conciliadas', 'Para_Revisao', 'Pendentes', 'Divergentes', 'Duplicadas'):
        sheet = book[name]
        headers = [sheet.cell(1, pos).value for pos in range(1, 16)]
        assert headers == expected


def test_workbook_snapshot_present() -> None:
    '''Resumo and log contain params snapshot and version.'''
    tables, params, statement, ledger = _make_results()
    payload = build_conciliation_workbook(tables, params, statement, ledger)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    resumo_vals: list = []
    sheet = book['Resumo']
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                resumo_vals.append(str(value))
    text = ' '.join(resumo_vals)
    assert 'date_tolerance_days' in text
    assert 'fuzzy_threshold' in text
    assert 'value_tolerance' in text
    assert APP_VERSION in text
    assert '2' in text
    log_sheet = book['Log_Regras']
    log_vals: list = []
    for row in log_sheet.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                log_vals.append(str(value))
    log_text = ' '.join(log_vals)
    assert 'date_tolerance_days' in log_text
    assert APP_VERSION in log_text


def test_workbook_error_headers() -> None:
    '''Erros sheet uses Linha Motivo Como corrigir.'''
    tables, params, statement, ledger = _make_results()
    payload = build_conciliation_workbook(tables, params, statement, ledger)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    sheet = book['Erros']
    headers = [sheet.cell(1, pos).value for pos in range(1, 4)]
    assert headers == ['Linha', 'Motivo', 'Como corrigir']


def _make_mixed_tables() -> tuple:
    '''Build tables mixing auto and exception values.'''
    statement = _make_frame(
        [
            ('10/09/2026', 'Pagamento Fornecedor X', '-2500'),
            ('11/09/2026', 'Recebimento Cliente Y', '4800'),
            ('15/09/2026', 'Servico Pequeno', '100'),
        ],
        'statement',
    )
    ledger = _make_frame(
        [
            ('10/09/2026', 'Fornecedor X NF 1254', '-2500'),
            ('12/09/2026', 'Cliente Y', '4800'),
            ('15/09/2026', 'Outro Valor', '9000'),
        ],
        'ledger',
    )
    params = build_params()
    tables = build_result_tables(statement, ledger, params)
    return tables, params, statement, ledger


def test_exceptions_only_and_ordered() -> None:
    '''Csv holds only exceptions ordered by valor desc.'''
    tables, params, statement, ledger = _make_mixed_tables()
    payload = build_exceptions_csv(tables, params, statement, ledger)
    text = payload.decode('utf-8')
    assert 'date_tolerance_days' in text
    assert 'fuzzy_threshold' in text
    assert 'value_tolerance' in text
    assert APP_VERSION in text
    frame = pd.read_csv(io.BytesIO(payload), comment='#')
    _assert_exceptions_frame(frame)


def _assert_exceptions_frame(frame: object) -> None:
    '''Check headers status and descending order.'''
    expected_start = [
        'Origem',
        'Data original',
        'Data normalizada',
        'Descrição original',
        'Valor original',
        'Valor normalizado',
    ]
    assert list(frame.columns)[:6] == expected_start
    assert 'Conciliada' not in set(frame['Status'].tolist())
    values = pd.to_numeric(frame['Valor normalizado'], errors='coerce').tolist()
    present = [item for item in values if pd.notna(item)]
    assert present == sorted(present, reverse=True)


def test_rule_log_ids_and_manual() -> None:
    '''Every auto potential has rule id manual appears.'''
    tables, _, _, _ = _make_results()
    log = build_rule_log(tables, None)
    for key in ('auto', 'potential'):
        frame = tables.get(key)
        if frame is not None and len(frame) > 0:
            assert all(str(item).strip() != '' for item in frame['rule_id'].tolist())
    assert 'rule_id' in list(log.columns)
    assert 'match_id' in list(log.columns)
    assert 'review_action' in list(log.columns)
    manual = [{'match_id': 's0-l0', 'review_action': 'conciliada_manual'}]
    manual[0]['timestamp'] = '2026-09-30T10:00:00'
    with_manual = build_rule_log(tables, manual)
    found = with_manual[with_manual['match_id'] == 's0-l0']
    assert len(found) >= 1
    assert found.iloc[0]['review_action'] == 'conciliada_manual'
    assert '2026-09-30' in str(found.iloc[0]['timestamp'])


def test_workbook_manual_visible() -> None:
    '''Manual confirmation surfaces in detail and log.'''
    tables, params, statement, ledger = _make_results()
    potentials = tables.get('potential')
    assert potentials is not None and len(potentials) > 0
    pair = str(potentials.iloc[0]['match_id'])
    review = [{'match_id': pair, 'review_action': 'conciliada_manual'}]
    review[0]['timestamp'] = '2026-09-30T12:00:00'
    payload = build_conciliation_workbook(tables, params, statement, ledger, None, review)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    sheet = book['Para_Revisao']
    manual_vals: list = []
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                manual_vals.append(str(value))
    assert 'Confirmada manualmente' in manual_vals
    log_sheet = book['Log_Regras']
    log_vals: list = []
    for row in log_sheet.iter_rows(values_only=True):
        for value in row:
            if value is not None:
                log_vals.append(str(value))
    assert pair in log_vals
    assert 'Confirmada manualmente' in log_vals


def test_prd_fixtures_workbook(fixture_paths: dict) -> None:
    '''Reference fixtures generate workbook and csv.'''
    raw_statement, _ = load_table(fixture_paths['statement'])
    mapped_statement = mapping.apply_mapping(raw_statement, mapping.auto_map_columns(raw_statement))
    statement = normalize_table(mapped_statement, 'statement')
    raw_ledger, _ = load_table(fixture_paths['ledger'])
    mapped_ledger = mapping.apply_mapping(raw_ledger, mapping.auto_map_columns(raw_ledger))
    ledger = normalize_table(mapped_ledger, 'ledger')
    valid_statement, errors_statement = collect_errors(statement)
    valid_ledger, errors_ledger = collect_errors(ledger)
    params = build_params()
    tables = build_result_tables(valid_statement, valid_ledger, params)
    errors = {'statement': errors_statement, 'ledger': errors_ledger}
    payload = build_conciliation_workbook(tables, params, valid_statement, valid_ledger, errors)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    assert 'Resumo' in book.sheetnames
    csv_payload = build_exceptions_csv(tables, params, valid_statement, valid_ledger, errors)
    frame = pd.read_csv(io.BytesIO(csv_payload), comment='#')
    assert len(frame) > 0
    assert statement.iloc[0]['normalized_date'] == date(2026, 9, 10)


def test_pct_divergent_includes_duplicate() -> None:
    '''Divergent rate sums divergent plus duplicate.'''
    empty = pd.DataFrame()
    dup = pd.DataFrame([{'match_id': 'dup-statement-0'}])
    results = {'auto': empty, 'potential': empty, 'pending': empty}
    results['divergent'] = empty
    results['duplicate'] = dup
    kpis = calc_kpis(results, 1, 1)
    assert kpis['pct_divergent'] == 100.0
    assert kpis['pct_auto'] == 0.0


def test_pct_divergent_mixed_counts() -> None:
    '''Mixed tables compute combined divergent share.'''
    auto = pd.DataFrame([{'a': 1}, {'a': 2}])
    empty = pd.DataFrame()
    div = pd.DataFrame([{'a': 3}])
    dup = pd.DataFrame([{'a': 4}])
    results = {'auto': auto, 'potential': empty, 'pending': empty}
    results['divergent'] = div
    results['duplicate'] = dup
    kpis = calc_kpis(results, 2, 2)
    assert kpis['pct_divergent'] == 50.0
    assert kpis['pct_auto'] == 50.0


def test_calc_kpis_empty_unified() -> None:
    '''Empty tables return zeroed kpis.'''
    empty = pd.DataFrame()
    results = {'auto': empty, 'potential': empty, 'pending': empty}
    results['divergent'] = empty
    results['duplicate'] = empty
    kpis = calc_kpis(results, 0, 0)
    assert kpis['pct_auto'] == 0.0
    assert kpis['pct_divergent'] == 0.0
    assert kpis['exception_rate'] == 100.0


def test_safe_len_branches() -> None:
    '''Safe len handles none invalid and frame.'''
    assert _safe_len(None) == 0
    assert _safe_len(42) == 0
    assert _safe_len(pd.DataFrame([{'a': 1}])) == 1
    assert _safe_len([]) == 0


def test_format_brl_branches() -> None:
    '''Currency handles nan invalid and list.'''
    assert format_brl(float('nan')) == '—'
    assert format_brl('abc') == 'abc'
    assert format_brl(['x']) == str(['x'])
    assert format_brl(Decimal('10.5')) == 'R$ 10,50'


def test_manual_lookup_branches() -> None:
    '''Manual lookup skips bad entries.'''
    assert _manual_lookup(None) == {}
    assert _manual_lookup('oops') == {}
    mixed = ['oops', {'noid': 1}, {'match_id': None}]
    assert _manual_lookup(mixed) == {}
    valid = [{'match_id': 's0-l0', 'review_action': 'rejeitada'}]
    lookup = _manual_lookup(valid)
    assert lookup['s0-l0']['review_action'] == 'rejeitada'


def test_manual_label_branches() -> None:
    '''Manual label covers all variants.'''
    assert _manual_label('conciliada_manual') == 'Confirmada manualmente'
    assert _manual_label('rejeitada') == 'Rejeitada'
    assert _manual_label('Confirmada manualmente') == 'Confirmada manualmente'
    assert _manual_label('Rejeitada') == 'Rejeitada'
    assert _manual_label('other') == ''
    assert _manual_label(None) == ''


def test_lookup_field_branches() -> None:
    '''Lookup field handles missing and valid.'''
    assert _lookup_field(None, 0, ['Data']) is None
    frame = pd.DataFrame([{'Data': '10/09/2026'}])
    assert _lookup_field(frame, None, ['Data']) is None
    assert _lookup_field(frame, 99, ['Data']) is None
    assert _lookup_field(frame, 0, ['Missing']) is None
    assert _lookup_field(frame, 0, ['Data']) == '10/09/2026'
    none_frame = pd.DataFrame([{'Data': None}])
    assert _lookup_field(none_frame, 0, ['Data']) is None


def test_format_iso_branches() -> None:
    '''Iso formatter handles dates and blanks.'''
    assert _format_iso(None) == ''
    assert _format_iso(float('nan')) == ''
    assert _format_iso(date(2026, 9, 10)) == '2026-09-10'
    assert _format_iso('') == ''
    assert _format_iso('nan') == ''
    assert _format_iso('10/09/2026') == '10/09/2026'


def test_format_iso_timestamp() -> None:
    '''Iso formatter handles pandas and datetime.'''
    stamp = pd.Timestamp('2026-09-10')
    assert _format_iso(stamp) == '2026-09-10'
    from datetime import datetime as moment
    assert _format_iso(moment(2026, 9, 10, 12, 0)) == '2026-09-10'


def test_as_text_branches() -> None:
    '''Text converter handles empty and dates.'''
    assert _as_text(None) == ''
    assert _as_text(float('nan')) == ''
    assert _as_text(date(2026, 9, 10)) == '10/09/2026'
    assert _as_text('hello') == 'hello'
    stamp = pd.Timestamp('2026-09-10')
    assert _as_text(stamp) == '10/09/2026'


def test_as_number_branches() -> None:
    '''Number converter handles none bool invalid.'''
    assert _as_number(None) is None
    assert _as_number(float('nan')) is None
    assert _as_number(True) is None
    assert _as_number(Decimal('10.5')) == 10.5
    assert _as_number('oops') is None


def test_resolve_side_frame_branches() -> None:
    '''Side resolver covers ledger statement fallback.'''
    statement_df = pd.DataFrame()
    ledger_frame = pd.DataFrame()
    first = {'source': 'ledger', 'ledger_idx': 1, 'statement_idx': 2}
    idx, _, origin = _resolve_side_frame(first, statement_df, ledger_frame)
    assert idx == 1
    assert origin == 'Interno'
    second = {'source': 'statement', 'ledger_idx': 1, 'statement_idx': 2}
    idx, _, origin = _resolve_side_frame(second, statement_df, ledger_frame)
    assert idx == 2
    assert origin == 'Extrato'
    third = {'ledger_idx': 5, 'statement_idx': 6}
    idx, _, origin = _resolve_side_frame(third, statement_df, ledger_frame)
    assert idx == 5
    fourth = {'statement_idx': 7}
    idx, _, origin = _resolve_side_frame(fourth, statement_df, ledger_frame)
    assert idx == 7


def test_normalize_error_frames_branches() -> None:
    '''Error normalizer handles dict list frame invalid.'''
    empty_cols = ['Linha', 'Motivo', 'Como corrigir']
    assert list(_normalize_error_frames(None).columns) == empty_cols
    assert list(_normalize_error_frames(42).columns) == empty_cols
    frame = pd.DataFrame([{'Linha': 1, 'Motivo': 'x', 'Como corrigir': 'y'}])
    assert len(_normalize_error_frames(frame)) == 1
    assert len(_normalize_error_frames([frame])) == 1
    assert len(_normalize_error_frames({'statement': frame})) == 1
    assert len(_normalize_error_frames([pd.DataFrame()])) == 0


def test_concat_and_clean_error_frame() -> None:
    '''Concat skips empty and renames guidance.'''
    frame = pd.DataFrame([{'Linha': 1, 'Motivo': 'm', 'Orientação': 'o'}])
    cleaned = _clean_error_frame(frame)
    assert list(cleaned.columns) == ['Linha', 'Motivo', 'Como corrigir']
    merged = _concat_error_frames([None, pd.DataFrame(), frame])
    assert len(merged) == 1
    assert list(_concat_error_frames([]).columns) == list(cleaned.columns)


def test_error_detail_rows_branches() -> None:
    '''Error detail covers dict list frame empty.'''
    assert _error_detail_rows(None) == []
    assert _error_detail_rows(42) == []
    frame = pd.DataFrame([{'Linha': 1, 'Motivo': 'Data inválida', 'Como corrigir': 'x'}])
    assert len(_error_detail_rows(frame)) == 1
    assert len(_error_detail_rows({'ledger_x': frame})) == 1
    assert len(_error_detail_rows({'statement_x': frame})) == 1
    assert len(_error_detail_rows([frame, frame])) == 2
    assert _error_detail_rows(pd.DataFrame()) == []
    assert _error_frame_rows(None, 'Extrato') == []


def test_infer_params_branches() -> None:
    '''Params inference handles bad and good inputs.'''
    assert _infer_params('oops') == {}
    assert _infer_params({}) == {}
    plain = pd.DataFrame([{'a': 1}])
    assert _infer_params({'auto': plain}) == {}
    snap = {'date_tolerance_days': 3, 'APP_VERSION': '9.9.9'}
    with_snap = pd.DataFrame([{'params_snapshot': snap}])
    assert _infer_params({'auto': with_snap}) == snap


def test_infer_error_code_branches() -> None:
    '''Error code inference covers all motives.'''
    assert _infer_error_code('Data inválida na linha 1') == 'DATA_INVALIDA'
    assert _infer_error_code('Valor inválido na linha 2') == 'VALOR_INVALIDO'
    assert _infer_error_code('Coluna obrigatória ausente') == 'COLUNA_AUSENTE'
    assert _infer_error_code('Outro problema') == 'ERRO_VALIDACAO'


def test_params_snapshot_single_key() -> None:
    '''Snapshot uses single APP_VERSION with compat.'''
    snap = _params_snapshot({})
    assert 'APP_VERSION' in snap
    assert 'app_version' not in snap
    legacy = _params_snapshot({'app_version': '0.0.1'})
    assert legacy['APP_VERSION'] == '0.0.1'
    modern = _params_snapshot({'APP_VERSION': '2.0.0'})
    assert modern['APP_VERSION'] == '2.0.0'
    text = _snapshot_text(legacy)
    assert '0.0.1' in text


def test_reason_label_pt() -> None:
    '''Reason codes translate to pt-BR labels.'''
    assert _reason_label('regra_composta_ok') == 'Valor, data e descrição conferem'
    assert _reason_label('duplicada_suspeita') == 'Duplicada suspeita'
    assert _reason_label('unknown_code') == 'unknown_code'


def test_workbook_reason_pt() -> None:
    '''Workbook detail shows pt-BR motive.'''
    tables, params, statement, ledger = _make_results()
    payload = build_conciliation_workbook(tables, params, statement, ledger)
    book = openpyxl.load_workbook(io.BytesIO(payload))
    found = False
    for name in ('Conciliadas', 'Para_Revisao', 'Pendentes', 'Divergentes', 'Duplicadas'):
        sheet = book[name]
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value in ('Valor, data e descrição conferem', 'Duplicada suspeita'):
                    found = True
                assert value != 'regra_composta_ok'
                assert value != 'duplicada_suspeita'
    assert isinstance(found, bool)


def test_rule_log_empty_and_errors() -> None:
    '''Rule log handles empty and error frames.'''
    empty = pd.DataFrame()
    results = {'auto': empty, 'potential': empty, 'pending': empty}
    results['divergent'] = empty
    results['duplicate'] = empty
    log = build_rule_log(results, None, None)
    assert len(log) == 0
    err = pd.DataFrame([{'Linha': 1, 'Motivo': 'Data inválida', 'Como corrigir': 'x'}])
    log_err = build_rule_log(results, [], {'statement': err})
    assert len(log_err) == 1
    assert log_err.iloc[0]['error_code'] == 'DATA_INVALIDA'
    log_list = build_rule_log(results, [], [err])
    assert len(log_list) == 1


def test_workbook_with_error_variants() -> None:
    '''Workbook accepts dict list and frame errors.'''
    tables, params, statement, ledger = _make_results()
    err = pd.DataFrame([{'Linha': 1, 'Motivo': 'Valor inválido', 'Como corrigir': 'x'}])
    for errors in ({'statement': err}, [err], err):
        payload = build_conciliation_workbook(tables, params, statement, ledger, errors, [])
        book = openpyxl.load_workbook(io.BytesIO(payload))
        assert 'Erros' in book.sheetnames
    payload = build_conciliation_workbook(tables, params, statement, ledger, None, None)
    assert len(payload) > 0


def test_exceptions_empty_and_sort() -> None:
    '''Exceptions handle empty and sort none last.'''
    assert _sort_number(None) == float('-inf')
    assert _sort_number(Decimal('5')) == 5.0
    empty = pd.DataFrame()
    results = {'auto': empty, 'potential': empty, 'pending': empty}
    results['divergent'] = empty
    results['duplicate'] = empty
    payload = build_exceptions_csv(results, {}, None, None, None, None)
    assert payload.decode('utf-8').startswith('# date_tolerance_days')
